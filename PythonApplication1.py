import tkinter as tk
from tkinter import ttk, filedialog
from tkinter.scrolledtext import ScrolledText
from tkinter import PhotoImage
import os

from docx import Document
from openpyxl import load_workbook

class DiceScribeApp:
    def __init__(self, root):
        self.root = root
        self.root.title("DiceScribe")
        self.root.geometry("900x600")

        self.page_control = ttk.Notebook(root)
        self.page_control.pack(expand=True, fill='both')

        self.home_page = ttk.Frame(self.page_control)
        self.directory_page = ttk.Frame(self.page_control)
        self.editor_page = ttk.Frame(self.page_control)
        self.manage_notes_page = ttk.Frame(self.page_control) 

        self.page_control.add(self.home_page, text="\U0001F3E0 Home")
        self.page_control.add(self.directory_page, text="\U0001F4D8 Directory")
        self.page_control.add(self.editor_page, text="\U0001F5C2\uFE0F Editor")
        self.page_control.add(self.manage_notes_page, text="Manage Notes")


        self.pages = {}
        self.history = []

        self.build_home()
        self.build_directory()
        self.build_editor()
        self.build_manage_notes()

    def build_home(self):
        top_frame = ttk.Frame(self.home_page)
        top_frame.pack(fill='x', pady=10)

        btn_dir = ttk.Button(top_frame, text="\U0001F4D8 Go to Directory", command=lambda: self.page_control.select(self.directory_page))
        btn_dir.pack(side='left', padx=5)

        btn_qs = ttk.Button(top_frame, text="\U0001F5C2\uFE0F Go to Editor", command=lambda: self.page_control.select(self.editor_page))
        btn_qs.pack(side='left', padx=5)

                
                
            
                

        desc = ("Welcome to DiceScribe - your all-in-one template and file editor!\n\n"
                "DiceScribe is a simple, intuitive platform that allows you to import and edit your gameplay files,\n"
                "and utilize one of our many customizable templates to speed up your note-taking and campaign management.\n\n"
                "To get started, just head to the Directory and import your first file.")

        label = ttk.Label(self.home_page, text=desc, font=("Arial", 12), wraplength=800, justify="center")
        label.pack(pady=20)

    def build_directory(self):
        container = ttk.Frame(self.directory_page, padding=10)
        container.pack(expand=True, fill='both')

        label = ttk.Label(container, text="Directory", font=("Arial", 16, "bold"))
        label.pack(pady=(0, 10))

        self.scenes = [
            ("Session 12: Betrayal at the Keep", "The party is betrayed by a trusted ally inside the fortress."),
            ("Session 13: The Oracle's Warning", "The oracle delivers a cryptic warning about an approaching threat."),
            ("Session 14: Final Battle Under the Moon", "The final showdown takes place beneath a blood-red moon."),
        ]

        self.scene_button_container = ttk.Frame(container)
        self.scene_button_container.pack(expand=True, fill='both')

        for title, summary in self.scenes:
            self.create_scene_button(title, summary)

        add_button = ttk.Button(container, text="Add Note from File", command=self.add_note_from_file)
        add_button.pack(pady=10)

    def create_scene_button(self, title, summary):
        btn = ttk.Button(self.scene_button_container, text=title, command=lambda: self.open_scene_in_editor(title, summary))
        btn.pack(fill='x', pady=4)

    def add_note_from_file(self):
        file_path = filedialog.askopenfilename(
            title="Select a Note File",
            filetypes=[
                ("Text Files", "*.txt"),
                ("Word Documents", "*.docx"),
                ("Excel Files", "*.xlsx")
            ]
        )
        if file_path:
            title = os.path.basename(file_path)
            ext = os.path.splitext(file_path)[1].lower()
            content = ""

            try:
                if ext == ".txt":
                    with open(file_path, "r", encoding="utf-8") as file:
                        content = file.read()
                elif ext == ".docx":
                    doc = Document(file_path)
                    content = "\n".join([para.text for para in doc.paragraphs])
                elif ext == ".xlsx":
                    wb = load_workbook(file_path, data_only=True)
                    sheet = wb.active
                    content = "\n".join([
                        "\t".join([str(cell.value) if cell.value is not None else "" for cell in row])
                        for row in sheet.iter_rows()
                    ])
                else:
                    print("Unsupported file type.")
                    return

                self.scenes.append((title, content))
                self.create_scene_button(title, content)
            except Exception as e:
                print(f"Error reading file: {e}")

    def open_scene_in_editor(self, title, content):

        if not hasattr(self, 'notebook'):
            self.build_editor()

        if title in self.pages:
            self.select_tab(title)
        else:
            self.add_editor_page(title, content)
            self.select_tab(title)
        self.page_control.select(self.editor_page)

    def build_editor(self):
        label = ttk.Label(self.editor_page, text="Editor", font=("Arial", 16, "bold"))
        label.pack(pady=(0, 10))

        self.notebook = ttk.Notebook(self.editor_page)
        self.notebook.pack(expand=True, fill='both')

        self.add_editor_page("Manage Notes", "Use this space to track, organize, and manage your note structure.")
        self.add_editor_page("Notes", "This is where you can write down your notes anytime.")
        self.add_editor_page("Character Sheet", "This is where all your character info is stored.")
        self.add_editor_page("Campaign", "This is where you can keep all your notes on your campaigns.")

        self.notebook.bind("<Button-1>", self.track_page)
        self.notebook.bind_all("<Control-Tab>", self.next_tab)
        self.notebook.bind_all("<Control-Shift-Tab>", self.prev_tab)

    def add_editor_page(self, title, content):
        frame = ttk.Frame(self.notebook)
        text_widget = ScrolledText(frame, wrap="word", font=("Arial", 11))
        text_widget.insert("1.0", content)
        text_widget.pack(expand=True, fill='both')

        save_button = ttk.Button(frame, text="Save", command=lambda: self.save_content(title, text_widget))
        save_button.pack(pady=5)

        self.notebook.add(frame, text=title)
        self.pages[title] = text_widget

    def build_manage_notes(self):
        frame = ttk.Frame(self.manage_notes_page, padding=10)
        frame.pack(fill='both', expand=True)

        label = ttk.Label(frame, text="Manage Notes", font=("Arial", 16, "bold"))
        label.pack(pady=(0, 10))

        btn = ttk.Button(frame, text="Organize Notes", command=self.manage_note_structure)
        btn.pack(pady=10)

    def get_folder_path(self, prompt, base_path):
        name = tk.simpledialog.askstring("Input", f"Enter {prompt} name:")
        if name:
            path = os.path.join(base_path, name)
            os.makedirs(path, exist_ok=True)
            return path
        return None

    def manage_note_structure(self):
        base = "C://Program Files/DiceScribe/GameNotes"
        game_path = self.get_folder_path("Game", base)
        if not game_path: return

        campaign_path = self.get_folder_path("Campaign", game_path)
        if not campaign_path: return

        category_path = self.get_folder_path("Category", campaign_path)
        if not category_path: return

        note_type_path = self.get_folder_path("Note type", category_path)
        if not note_type_path: return

        move_files = tk.messagebox.askyesno("Move Files", "Would you like to move notes to this folder?")
        if move_files:
            file_paths = filedialog.askopenfilenames(title="Select Notes to Move")
            for path in file_paths:
                try:
                    os.rename(path, os.path.join(note_type_path, os.path.basename(path)))
                except Exception as e:
                    print(f"Error moving file {path}: {e}")

    def save_content(self, title, text_widget):
        file_path = filedialog.asksaveasfilename(defaultextension=".txt", initialfile=title,
                                                 filetypes=[("Text Files", "*.txt")])
        if file_path:
            try:
                content = text_widget.get("1.0", tk.END)
                with open(file_path, "w", encoding="utf-8") as file:
                    file.write(content)
            except Exception as e:
                print(f"Error saving file: {e}")

    def track_page(self, event=None):
        current_tab = self.notebook.tab(self.notebook.select(), "text")
        if current_tab in self.history:
            self.history.remove(current_tab)
        self.history.insert(0, current_tab)

    def next_tab(self, event=None):
        if self.history:
            current_tab = self.notebook.tab(self.notebook.select(), "text")
            try:
                index = self.history.index(current_tab)
                next_index = (index + 1) % len(self.history)
                self.select_tab(self.history[next_index])
            except ValueError:
                pass

    def prev_tab(self, event=None):
        if self.history:
            current_tab = self.notebook.tab(self.notebook.select(), "text")
            try:
                index = self.history.index(current_tab)
                prev_index = (index - 1) % len(self.history)
                self.select_tab(self.history[prev_index])
            except ValueError:
                pass

    def select_tab(self, title):
        for tab_id in self.notebook.tabs():
            if self.notebook.tab(tab_id, "text") == title:
                self.notebook.select(tab_id)
                self.track_page()
                break

if __name__ == "__main__":
    root = tk.Tk()
    app = DiceScribeApp(root)
    root.mainloop()