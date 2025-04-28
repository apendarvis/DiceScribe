import tkinter as tk
from tkinter import ttk
from tkinter import filedialog
from tkinter import simpledialog
import subprocess
import os
import random
import shutil
from tkinter import messagebox
from docx import Document
from PIL import Image, ImageTk
from openpyxl import load_workbook
from tkinter.scrolledtext import ScrolledText

libreoffice_path = r"C:\Program Files\LibreOffice\program"

def open_with_LibreOffice():
    file_path = filedialog.askopenfilename(title="Choose a file to open")
    if file_path:
        ext = os.path.splitext(file_path)[1].lower()
        if ext in ['.ods', '.xls', '.xlsx']:
            subprocess.Popen(["C:\Program Files\LibreOffice\program\scalc.exe", file_path])
        elif ext in ['.odt', '.doc', '.docx']:
            subprocess.Popen(["C:\Program Files\LibreOffice\program\swriter.exe", file_path])
        else:
            subprocess.Popen(["C:\Program Files\LibreOffice\program\soffice.exe", file_path])  # fallback



def open_folder_with_LibreOffice():
    folder_path = filedialog.askdirectory(title="Choose a folder to open files from")
    if folder_path:
        for filename in os.listdir(folder_path):
            file_path = os.path.join(folder_path, filename)
            if os.path.isfile(file_path):
                ext = os.path.splitext(file_path)[1].lower()
                if ext in ['.ods', '.xls', '.xlsx']:
                    subprocess.Popen([os.path.join(libreoffice_path, "scalc.exe"), file_path])
                elif ext in ['.odt', '.doc', '.docx']:
                    subprocess.Popen([os.path.join(libreoffice_path, "swriter.exe"), file_path])
               
                else:
                    subprocess.Popen([os.path.join(libreoffice_path, "soffice.exe"), file_path])

def copy_and_rename_and_open_with_LibreOffice():
    # Ask the user to choose the file to copy
    source_file_path = filedialog.askopenfilename(title="Choose a file to copy")
    if source_file_path:
        # Ask for the destination folder to copy the file to
        dest_folder = filedialog.askdirectory(title="Choose a folder to copy the file to")
        if dest_folder:
            # Ask the user to provide a new name for the file
            new_file_name = simpledialog.askstring("Rename File", "Enter a new name for the file (with extension):")
            if new_file_name:
                try:
                    # Ensure the new file name has the correct extension
                    ext = os.path.splitext(source_file_path)[1]
                    if not new_file_name.endswith(ext):
                        new_file_name += ext
                    
                    # Create the full destination path for the renamed file
                    dest_file_path = os.path.join(dest_folder, new_file_name)
                    
                    # Copy the file to the new location with the new name
                    shutil.copy(source_file_path, dest_file_path)
                    print(f"File copied to {dest_file_path}")
                    
                    # Now open the copied (and renamed) file with LibreOffice
                    ext = os.path.splitext(dest_file_path)[1].lower()
                    if ext in ['.ods', '.xls', '.xlsx']:
                        subprocess.Popen([os.path.join(libreoffice_path, "scalc.exe"), dest_file_path])
                    elif ext in ['.odt', '.doc', '.docx']:
            
                        subprocess.Popen([os.path.join(libreoffice_path, "swriter.exe"), dest_file_path])
                    else:
                        subprocess.Popen([os.path.join(libreoffice_path, "soffice.exe"), dest_file_path])  # fallback
                    
                    messagebox.showinfo("Success", f"File copied and opened with LibreOffice.")
                except Exception as e:
                    messagebox.showerror("Error", f"Failed to copy the file: {e}")
            else:
                messagebox.showwarning("No Name", "You must provide a new name for the file.")

def make_new_campaign():
    # Prompt user to select the location for the new campaign folder
    campaign_location = filedialog.askdirectory(title="Choose a location for the new campaign folder")
    if campaign_location:
        # Ask the user for the new campaign folder name
        campaign_name = simpledialog.askstring("Campaign Name", "Enter the name for the new campaign:")
        if campaign_name:
            try:
                # Create the new campaign folder path
                new_campaign_path = os.path.join(campaign_location, campaign_name)
                
                # Check if the folder already exists
                if not os.path.exists(new_campaign_path):
                    os.makedirs(new_campaign_path)
                    
                    # Create subfolders: DM Only, Players, General
                    os.makedirs(os.path.join(new_campaign_path, "DM Only"))
                    os.makedirs(os.path.join(new_campaign_path, "Players"))
                    os.makedirs(os.path.join(new_campaign_path, "General"))
                    
                    messagebox.showinfo("Success", f"New campaign folder '{campaign_name}' created with subfolders.")
                else:
                    messagebox.showwarning("Folder Exists", "A folder with that name already exists.")
            except Exception as e:
                messagebox.showerror("Error", f"Failed to create the campaign folder: {e}")
        else:
            messagebox.showwarning("No Name", "You must provide a name for the campaign folder.")
            

class QuickSwitchApp:
    def __init__(self, root):
        self.root = root
        self.root.title("DiceScribe - QuickSwitch")
        
        self.notebook = ttk.Notebook(root)
        self.notebook.pack(expand=True, fill='both')
        
        self.pages = {}
        self.history = []
        
        # Sample pages
        self.add_page("Notes", "View and write notes.") 
        self.add_page("Character Sheets", "View and create Character Sheets.")
        self.add_page("Campaigns", "View and create Campiagns.")
        self.add_page("Damage Calculator" , "Calculate Damage And Saving Rolls.")
        
        self.notebook.bind("<Button-1>", self.track_page)
        self.notebook.bind_all("<Control-Tab>", self.next_tab)
        self.notebook.bind_all("<Control-Shift-Tab>", self.prev_tab)
        
    def add_page(self, title, text=""):
        frame = ttk.Frame(self.notebook)
        label = ttk.Label(frame, text=text, padding=10)
        label.pack(expand=True)
        if title == "Character Sheets":
            open_button = ttk.Button(frame, text="Open Character Sheet with LibreOffice", command=open_with_LibreOffice)
            open_button.pack(pady=10)
            copy_rename_open_button = ttk.Button(frame, text="Use a Character Sheet Template", command=copy_and_rename_and_open_with_LibreOffice)
            copy_rename_open_button.pack(pady=10)
        if title == "Notes":
            open_button = ttk.Button(frame, text="Open Notes with LibreOffice", command=open_with_LibreOffice)
            open_button.pack(pady=10)
        if title == "Campaigns":
            open_button = ttk.Button(frame, text="Open all Campaign files", command=open_folder_with_LibreOffice)
            open_button.pack(pady=10)
            new_campaign_button = ttk.Button(frame, text="Create New Campaign", command=make_new_campaign)
            new_campaign_button.pack(pady=10)
            
        
        
    def track_page(self, event=None):
        current_tab = self.notebook.tab(self.notebook.select(), "text")
        if current_tab in self.history:
            self.history.remove(current_tab)
        self.history.insert(0, current_tab)
        
    def next_tab(self, event=None):
        if self.history:
            current_tab = self.notebook.tab(self.notebook.select(), "text")
            try:
                index = self.history.index(current_tab)
                next_index = (index + 1) % len(self.history)
                self.select_tab(self.history[next_index])
            except ValueError:
                pass
                
    def prev_tab(self, event=None):
        if self.history:
            current_tab = self.notebook.tab(self.notebook.select(), "text")
            try:
                index = self.history.index(current_tab)
                prev_index = (index - 1) % len(self.history)
                self.select_tab(self.history[prev_index])
            except ValueError:
                pass
                
    def select_tab(self, title):
        for tab_id in self.notebook.tabs():
            if self.notebook.tab(tab_id, "text") == title:
                self.notebook.select(tab_id)
                self.track_page()
                break
            

class DiceScribeApp:
    def __init__(self, root):
        self.root = root
        self.root.title("DiceScribe")
        self.root.geometry("900x600")

        self.page_control = ttk.Notebook(root)
        self.page_control.pack(expand=True, fill='both')

        self.home_page = ttk.Frame(self.page_control)
        self.directory_page = ttk.Frame(self.page_control)
        self.editor_page = ttk.Frame(self.page_control)
        self.damage_page = ttk.Frame(self.page_control)

        self.page_control.add(self.home_page, text="\U0001F3E0 Home")
        self.page_control.add(self.directory_page, text="\U0001F4D8 Directory")
        self.page_control.add(self.editor_page, text="\U0001F5C2\uFE0F Editor")
        self.page_control.add(self.damage_page, text="\U0001F3AF Damage Calculator")

        self.pages = {}
        self.history = []

        self.build_home()
        self.build_directory()
        self.build_editor()
        self.build_damage_calc()
        

    def build_home(self):
        top_frame = ttk.Frame(self.home_page)
        top_frame.pack(fill='x', pady=10)

        btn_dir = ttk.Button(top_frame, text="\U0001F4D8 Go to Directory", command=lambda: self.page_control.select(self.directory_page))
        btn_dir.pack(side='left', padx=5)

        btn_qs = ttk.Button(top_frame, text="\U0001F5C2\uFE0F Go to Editor", command=lambda: self.page_control.select(self.editor_page))
        btn_qs.pack(side='left', padx=5)

                
                
            
                

        desc = ("Welcome to DiceScribe — your all-in-one template and file editor!\n\n"
                "DiceScribe is a simple, intuitive platform that allows you to import and edit your gameplay files,\n"
                "and utilize one of our many customizable templates to speed up your note-taking and campaign management.\n\n"
                "To get started, just head to the Directory and import your first file.")

        label = ttk.Label(self.home_page, text=desc, font=("Arial", 12), wraplength=800, justify="center")
        label.pack(pady=20)
        
    def build_damage_calc(self):
        container = ttk.Frame(self.damage_page, padding=10)
        container.pack(expand=True, fill='both')

        label = ttk.Label(container, text="Damage Calculator", font=("Arial", 16, "bold"))
        label.pack(pady=(0, 10))

        ttk.Label(container, text="Enter Number of Dice:").pack(pady=5)
        self.dice_count_entry = ttk.Entry(container, width=20)
        self.dice_count_entry.pack(pady=5)
        
        ttk.Label(container, text="Enter Sides of Die:").pack(pady=5)
        self.dice_sides_entry = ttk.Entry(container, width=20)
        self.dice_sides_entry.pack(pady=5)
        
        ttk.Label(container, text="Enter Modifier(+ , -):").pack(pady=5)
        self.modifier_entry = ttk.Entry(container, width=20)
        self.modifier_entry.pack(pady=5)

        calc_button = ttk.Button(container, text="Calculate Damage", command=self.calculate_damage)
        calc_button.pack(pady=10)

        self.result_label = ttk.Label(container, text="Result: ", font=("Arial", 12))
        self.result_label.pack(pady=20)

    def calculate_damage(self):
        try:
            #get the values 
            dice_count = int(self.dice_count_entry.get())
            dice_sides = int(self.dice_sides_entry.get())
            modifier = int(self.modifier_entry.get()) if self.modifier_entry.get() else 0
            
            
            dice_rolls = [random.randint(1, dice_sides) for _ in range(dice_count)]
            total_roll = sum(dice_rolls)

            # calculate the total damage
            total_damage = total_roll + modifier

            #display 
            self.result_label.config(text=f"Result: {total_damage} (You Rolled: {', '.join(map(str, dice_rolls))}, Modifier: {modifier})")
    
        except ValueError:
            self.result_label.config(text="Error: Please enter valid integers for all fields.")

        

    def build_directory(self):
        container = ttk.Frame(self.directory_page, padding=10)
        container.pack(expand=True, fill='both')

        label = ttk.Label(container, text="Directory", font=("Arial", 16, "bold"))
        label.pack(pady=(0, 10))

        self.scenes = [
            ("Session 12: Betrayal at the Keep", "The party is betrayed by a trusted ally inside the fortress."),
            ("Session 13: The Oracle's Warning", "The oracle delivers a cryptic warning about an approaching threat."),
            ("Session 14: Final Battle Under the Moon", "The final showdown takes place beneath a blood-red moon."),
        ]

        self.scene_button_container = ttk.Frame(container)
        self.scene_button_container.pack(expand=True, fill='both')

        for title, summary in self.scenes:
            self.create_scene_button(title, summary)

        add_button = ttk.Button(container, text="➕ Add Note from File", command=self.add_note_from_file)
        add_button.pack(pady=10)

    def create_scene_button(self, title, summary):
        btn = ttk.Button(self.scene_button_container, text=title, command=lambda: self.open_scene_in_editor(title, summary))
        btn.pack(fill='x', pady=4)

    def add_note_from_file(self):
        file_path = filedialog.askopenfilename(
            title="Select a Note File",
            filetypes=[
                ("Text Files", "*.txt"),
                ("Word Documents", "*.docx"),
                ("Excel Files", "*.xlsx"),
                ("Image Files", "*.jpg; *.jpeg; *.png; *.gif")
            ]
        )
        if file_path:
            title = os.path.basename(file_path)
            ext = os.path.splitext(file_path)[1].lower()
            content = ""

            try:
                if ext == ".txt":
                    with open(file_path, "r", encoding="utf-8") as file:
                        content = file.read()
                elif ext == ".docx":
                    doc = Document(file_path)
                    content = "\n".join([para.text for para in doc.paragraphs])
                elif ext == ".xlsx":
                    wb = load_workbook(file_path, data_only=True)
                    sheet = wb.active
                    content = "\n".join([
                        "\t".join([str(cell.value) if cell.value is not None else "" for cell in row])
                        for row in sheet.iter_rows()
                    ])
                elif ext in [".jpg", ".jpeg", ".png", ".gif"]:
                    self.add_image_page(title, file_path)
                    self.page_control.select(self.editor_page)
                else:
                    print("Unsupported file type.")
                    return

                self.scenes.append((title, content))
                self.create_scene_button(title, content)
            except Exception as e:
                print(f"Error reading file: {e}")
    def add_image_page(self, title, image_path):
        try:
            img = Image.open(image_path)
            img.thumbnail((800, 600))
            photo = ImageTk.PhotoImage(img)

            frame = ttk.Frame(self.notebook)
            label = tk.Label(frame, image=photo)
            label.image = photo
            label.pack(expand=True, fill = 'both')

            self.notebook.add(frame, text=title)
            self.pages[title] = label
            self.select_tab(title)
        except Exception as e:
            print(f"Error loading image: {e}")
    def open_scene_in_editor(self, title, content):
        if title in self.pages:
            self.select_tab(title)
        else:
            self.add_editor_page(title, content)
            self.select_tab(title)
        self.page_control.select(self.editor_page)
    
    def build_search_bar(self):
       
        search_frame = ttk.Frame(self.editor_page)
        search_frame.pack(pady=5, fill='x')

        ttk.Label(search_frame, text="Search:").pack(side="left", padx=(10, 5))
        self.search_var = tk.StringVar()
        search_entry = ttk.Entry(search_frame, textvariable=self.search_var, width=40)
        search_entry.pack(side="left", padx=(0, 10))
        search_entry.bind("<KeyRelease>", self.highlight_search)
        
    def highlight_search(self, event=None):
        # If the search bar is empty, remove all highlights and return
        if self.is_search_empty():
            current_tab_id = self.notebook.select()
            if not current_tab_id:
                return

            current_tab_title = self.notebook.tab(current_tab_id, "text")
            text_widget = self.pages.get(current_tab_title)

            if text_widget:
                text_widget.tag_remove("highlight", "1.0", tk.END)  # Remove all highlights
                return

        search_term = self.search_var.get()
        
        if len(search_term.strip()) < 4:
            # Also clear existing highlights
            current_tab_id = self.notebook.select()
            if current_tab_id:
                current_tab_title = self.notebook.tab(current_tab_id, "text")
                text_widget = self.pages.get(current_tab_title)
                if text_widget:
                    text_widget.tag_remove("highlight", "1.0", tk.END)
            return
        
        current_tab_id = self.notebook.select()
        if not current_tab_id:
            return

        current_tab_title = self.notebook.tab(current_tab_id, "text")
        text_widget = self.pages.get(current_tab_title)

        if not text_widget:
            print(f"No text widget found for tab {current_tab_title}.")
            return

        text_widget.tag_remove("highlight", "1.0", tk.END)
    
        startHighlight = "1.0"  #starts at line 1 character 0
        while True:
            #searches text widget 
            startHighlight = text_widget.search(search_term, startHighlight, nocase=True, stopindex=tk.END)
            if not startHighlight:
                break
            # finds start of found term and adds search term length to highlight full statement
            end = f"{startHighlight}+{len(search_term)}c"
            text_widget.tag_add("highlight", startHighlight, end)
            startHighlight = end

        # highlight style (yellow)
        text_widget.tag_config("highlight", background="yellow", foreground="black")
        
    def is_search_empty(self):
        return self.search_var.get().strip() == ""

    def build_editor(self):
        label = ttk.Label(self.editor_page, text="Editor", font=("Arial", 16, "bold"))
        label.pack(pady=(0, 10))
        
        self.build_search_bar()

        self.notebook = ttk.Notebook(self.editor_page)
        self.notebook.pack(expand=True, fill='both')

        self.add_editor_page("Notes Template", 
"""
## Notes

Welcome to your **Notes** page! Here you can quickly jot down any ideas, reminders, or thoughts you have during your game. This is where you'll organize and keep track of your campaign progress, character insights, or anything else that needs to be remembered. 

- **Add New Note**: Go to the Directory to import your already made notes.
- **Search Notes**: Use the search bar to find specific notes.
- **Edit Notes**: Modify any notes by clicking on the spot you want to modify and start editing.

Stay organized and keep your thoughts in one place!
""")
        self.add_editor_page("Character Sheet Template",
        """
## Character Sheet

This is where you can write down all your character information is stored. You can track your character's stats, abilities, inventory, and background details. The character sheet is designed to be a central hub for everything related to your character.

- Character Info: Name, class, race, etc.
- Abilities: Track your character's strength, dexterity, and other attributes.
- Inventory: Keep a detailed list of your character's equipment and items.

Use this page to build your character and keep their story consistent!
""")
        self.add_editor_page("Campaign Template", 
"""
## Campaign

This is the heart of your campaign. Keep all your notes here for each adventure, NPCs, plot twists, and locations. Stay organized with a clear overview of your campaign progress.

- Campaign Overview: Quick summary of the current campaign.
- NPCs: List all non-player characters, their backgrounds, and motivations.
- Locations: Keep track of important locations your party visits.

All your campaign essentials in one place!
""")

        self.notebook.bind("<Button-1>", self.track_page)
        self.notebook.bind_all("<Control-Tab>", self.next_tab)
        self.notebook.bind_all("<Control-Shift-Tab>", self.prev_tab)

    def add_editor_page(self, title, content):
        frame = ttk.Frame(self.notebook)
        text_widget = ScrolledText(frame, wrap="word", font=("Arial", 11))
        text_widget.insert("1.0", content)
        text_widget.pack(expand=True, fill='both')

        save_button = ttk.Button(frame, text="💾 Save", command=lambda: self.save_content(title, text_widget))
        save_button.pack(pady=5)

        self.notebook.add(frame, text=title)
        self.pages[title] = text_widget

    def save_content(self, title, text_widget):
        file_path = filedialog.asksaveasfilename(defaultextension=".txt", initialfile=title,
                                                 filetypes=[("Text Files", "*.txt")])
        if file_path:
            try:
                content = text_widget.get("1.0", tk.END)
                with open(file_path, "w", encoding="utf-8") as file:
                    file.write(content)
            except Exception as e:
                print(f"Error saving file: {e}")
                

    def track_page(self, event=None):
        current_tab = self.notebook.tab(self.notebook.select(), "text")
        if current_tab in self.history:
            self.history.remove(current_tab)
        self.history.insert(0, current_tab)

    def next_tab(self, event=None):
        if self.history:
            current_tab = self.notebook.tab(self.notebook.select(), "text")
            try:
                index = self.history.index(current_tab)
                next_index = (index + 1) % len(self.history)
                self.select_tab(self.history[next_index])
            except ValueError:
                pass

    def prev_tab(self, event=None):
        if self.history:
            current_tab = self.notebook.tab(self.notebook.select(), "text")
            try:
                index = self.history.index(current_tab)
                prev_index = (index - 1) % len(self.history)
                self.select_tab(self.history[prev_index])
            except ValueError:
                pass

    def select_tab(self, title):
        for tab_id in self.notebook.tabs():
            if self.notebook.tab(tab_id, "text") == title:
                self.notebook.select(tab_id)
                self.track_page()
                break

if __name__ == "__main__":
    root = tk.Tk()
    app = DiceScribeApp(root)
    root.mainloop()
